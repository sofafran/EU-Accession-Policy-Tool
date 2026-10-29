#!/usr/bin/env python3
"""Convert an original source file to a text file with page or section markers."""

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import nfc  # noqa: E402

__version__ = "1.0.0"
MIN_CHARS = 25  # a PDF page with fewer non-space characters than this is treated as scanned


def ocr_pdf_page(pdf, page_no, langs, dpi):
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        raise SystemExit("error: text recognition needs pdftoppm (poppler-utils) and tesseract")
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / "p"
        subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(page_no), "-l", str(page_no), "-png", str(pdf), str(base)],
                       check=True, capture_output=True)
        image = next(Path(tmp).glob("p*.png"))
        out = subprocess.run(["tesseract", str(image), "stdout", "-l", langs], check=True, capture_output=True, text=True)
        return out.stdout


def convert_pdf(path, langs, dpi, force_ocr):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts, ocr_pages = [], []
    for i, page in enumerate(reader.pages, start=1):
        text = "" if force_ocr else (page.extract_text() or "")
        if len("".join(text.split())) < MIN_CHARS:
            text = ocr_pdf_page(path, i, langs, dpi)
            ocr_pages.append(i)
        parts.append(f"=== PAGE {i} ===\n{text.strip()}\n")
    return "\n".join(parts), ocr_pages, len(reader.pages)


def convert_html(path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(Path(path).read_bytes(), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form"]):
        tag.decompose()
    body = soup.body or soup
    lines = ["=== PAGE 1 ==="]
    for el in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "blockquote", "pre"]):
        if el.find(["p", "li", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6"]):
            continue  # take text from the innermost block only
        text = " ".join(el.get_text(" ").split())
        if not text:
            continue
        if el.name.startswith("h"):
            lines.append(f"=== SECTION: {text} ===")
        else:
            lines.append(text)
    return "\n".join(lines) + "\n", [], 1


def convert_docx(path):
    import docx
    d = docx.Document(str(path))
    lines = ["=== PAGE 1 ==="]
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        lines.append(f"=== SECTION: {text} ===" if style.startswith(("heading", "title")) else text)
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines) + "\n", [], 1


def convert_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = [" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
        parts.append(f"=== PAGE {i} ===\n=== SECTION: {ws.title} ===\n" + "\n".join(r for r in rows if r.strip(" |")) + "\n")
    return "\n".join(parts), [], len(wb.worksheets)


def main(argv=None):
    p = argparse.ArgumentParser(
        description="Convert a PDF, HTML, DOCX or XLSX original to UTF-8 text (NFC). Writes '=== PAGE n ===' before each page; "
                    "webpages and Word files have no pages, so they get one page marker and '=== SECTION: heading ===' lines for headings. "
                    "PDF pages with no usable text are run through text recognition, and the tool reports which pages.")
    p.add_argument("original", help="original file (never modified)")
    p.add_argument("output", help="text file to write")
    p.add_argument("--ocr-langs", default="eng+mkd+sqi", help="tesseract languages for scanned pages (default eng+mkd+sqi, as in the run record)")
    p.add_argument("--dpi", type=int, default=300, help="resolution for text recognition (default 300)")
    p.add_argument("--force-ocr", action="store_true", help="run text recognition on every PDF page")
    p.add_argument("--version", action="version", version=__version__)
    a = p.parse_args(argv)
    src = Path(a.original)
    ext = src.suffix.lower()
    if ext == ".pdf":
        text, ocr_pages, pages = convert_pdf(src, a.ocr_langs, a.dpi, a.force_ocr)
    elif ext in (".html", ".htm"):
        text, ocr_pages, pages = convert_html(src)
    elif ext == ".docx":
        text, ocr_pages, pages = convert_docx(src)
    elif ext == ".xlsx":
        text, ocr_pages, pages = convert_xlsx(src)
    else:
        p.error(f"unsupported format: {ext}")
    out = Path(a.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(nfc(text), encoding="utf-8")
    print(f"wrote {out}: {pages} page(s)")
    if ocr_pages:
        print(f"ocr: text recognition used on page(s) {', '.join(map(str, ocr_pages))} ({a.ocr_langs}); note 'ocr' in the register")
    return 0


if __name__ == "__main__":
    sys.exit(main())
